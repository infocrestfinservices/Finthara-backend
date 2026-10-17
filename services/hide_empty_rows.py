"""
hide_empty_rows.py

Hides the line items a business does not have — a row of nil across all sixty months
(e.g. "Power & fuel" for a software company) — so the schedule shows what the business
actually spends on.

Hidden, never deleted: formulas, totals and every reference from other sheets stay exactly
as they are (totals use SUM, which counts hidden rows; no template uses SUBTOTAL), and the
reader can unhide a row at any time. A note on the sheet says so.

Runs on the RECALCULATED workbook only — before recalculation the formula cells carry no
values, so "every month is zero" cannot be known. When there is no recalculated file the
caller does not call this, and the workbook goes out with every row showing.

Edits the .xlsx as XML (a `hidden="1"` on each row element plus the note cell), copying
every other zip entry verbatim — the same rule template_fill_service follows. Re-saving
through openpyxl would drop the cover page's shapes and parts of the chart styling.
openpyxl is used here only to READ values and formulas.

A row is hidden only when ALL of these hold:
  * it sits below the sheet's month-header row ("M1", "M2", ...) and has a label;
  * its label is not a total or sub-total;
  * it carries a value or formula in the month columns (a label-only row is a heading);
  * every month cell is zero or blank;
  * its input is zero or blank — the Assumptions cells its formula reads (escalation terms
    like "(1+Assumptions!C33)" excepted), or, when it reads no Assumptions cell, the cells
    on other sheets it reads. A row with no references at all (a constant 0) has no input.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from services.template_fill_service import _sheet_name_to_path, _set_cell_xml

logger = logging.getLogger(__name__)

# The sheets this applies to. Add a sheet name here to extend it — the rules above are
# generic (month header, totals, inputs). Never add the bank's fixed formats: Form II-VI,
# WC & CC-OD Limit, DSCR, Ratios, Repayment.
HIDE_SHEETS = ["Expenses"]

NOTE = "Rows with no values are hidden. Unhide to enter values."

_INPUT_SHEET = "Assumptions"
_PROTECTED = re.compile(r"\b(sub-?\s*totals?|totals?)\b", re.I)
_MONTH = re.compile(r"^M\d+$")
_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)"
_REF = re.compile(r"(" + _SHEET + r")!\$?([A-Z]{1,3})\$?(\d+)")
_GROWTH = re.compile(r"\(\s*1\s*\+\s*" + _SHEET + r"!\$?[A-Z]{1,3}\$?\d+\s*\)")


def apply(xlsx_bytes: bytes, sheets=None) -> bytes:
    """The workbook with its empty rows hidden. Never fails a download: on any problem
    the workbook is returned unchanged."""
    try:
        return _apply(xlsx_bytes, HIDE_SHEETS if sheets is None else sheets)
    except Exception:
        logger.warning("hide_empty_rows: skipped, workbook left unchanged", exc_info=True)
        return xlsx_bytes


def _apply(xlsx_bytes: bytes, sheets) -> bytes:
    values = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    formulas = load_workbook(io.BytesIO(xlsx_bytes))

    plan = {}                                   # sheet -> (rows to hide, note cell or None)
    for name in sheets:
        if name not in values.sheetnames:
            continue
        found = _empty_rows(values, formulas, name)
        if found:
            plan[name] = found
    if not plan:
        return xlsx_bytes

    zin = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    paths = _sheet_name_to_path(zin)
    modified = {}
    for name, (rows, note_cell) in plan.items():
        path = paths.get(name)
        if not path:
            continue
        xml = zin.read(path).decode("utf-8")
        hidden = []
        for r in rows:
            xml, ok = _hide_row(xml, r)
            if ok:
                hidden.append(r)
        if not hidden:
            continue
        if note_cell:
            xml = _write_note(xml, note_cell)
        modified[path] = xml.encode("utf-8")
        logger.info("hide_empty_rows: %s rows %s hidden", name, hidden)
    if not modified:
        return xlsx_bytes

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = modified.get(item.filename)
            if data is None:
                data = zin.read(item.filename)
            zi = zipfile.ZipInfo(item.filename, date_time=item.date_time)
            zi.compress_type = item.compress_type
            zi.external_attr = item.external_attr
            zi.internal_attr = item.internal_attr
            zi.create_system = item.create_system
            zout.writestr(zi, data)
    return buf.getvalue()


def _is_zero(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return abs(v) < 1e-9
    return False


def _empty_rows(values, formulas, name):
    """(rows to hide, cell for the note) for one sheet, or None."""
    wv, wf = values[name], formulas[name]
    header, months = None, []
    for r in range(1, min(wv.max_row, 10) + 1):
        cols = [c for c in range(1, wv.max_column + 1)
                if _MONTH.match(str(wv.cell(r, c).value or "").strip())]
        if cols:
            header, months = r, cols
            break
    if not header:
        return None

    rows = []
    for r in range(header + 1, wv.max_row + 1):
        label = wv.cell(r, 1).value
        if label is None or not str(label).strip():
            continue
        if _PROTECTED.search(str(label)):
            continue
        if all(wf.cell(r, c).value is None for c in months):
            continue                              # a heading, not a line item
        if not all(_is_zero(wv.cell(r, c).value) for c in months):
            continue
        if not _input_is_zero(values, name, wf.cell(r, months[0]).value):
            continue
        rows.append(r)
    if not rows:
        return None

    # The note goes in column A of the empty row nearest above the month header.
    note_cell = None
    for r in range(header - 1, 0, -1):
        if wv.cell(r, 1).value is None and wf.cell(r, 1).value is None:
            note_cell = f"A{r}"
            break
    return rows, note_cell


def _input_is_zero(values, sheet, formula) -> bool:
    if not isinstance(formula, str) or not formula.startswith("="):
        return True                               # a constant: no input behind it
    refs = []
    for sh, col, row in _REF.findall(_GROWTH.sub("", formula)):
        sh = sh[1:-1].replace("''", "'") if sh.startswith("'") else sh
        if sh != sheet:
            refs.append((sh, f"{col}{row}"))
    chosen = [x for x in refs if x[0] == _INPUT_SHEET] or refs
    for sh, cell in chosen:
        if sh not in values.sheetnames:
            return False                          # cannot read it: do not hide
        if not _is_zero(values[sh][cell].value):
            return False
    return True


def _hide_row(xml: str, r: int):
    pattern = re.compile(r'<((?:\w+:)?row)\b([^>]*?)(/?)>')

    def find(m):
        return re.search(r'\br="%d"' % r, m.group(2)) is not None

    for m in pattern.finditer(xml):
        if find(m):
            attrs = re.sub(r'\s+hidden="[^"]*"', "", m.group(2))
            new = f'<{m.group(1)}{attrs} hidden="1"{m.group(3)}>'
            return xml[:m.start()] + new + xml[m.end():], True
    return xml, False


def _write_note(xml: str, cell: str) -> str:
    """Put the note in `cell` (column A). Overwrites an existing empty cell element, keeping
    its style; otherwise inserts a new one as the row's first cell."""
    new_xml, ok = _set_cell_xml(xml, cell, "str", NOTE)
    if ok:
        return new_xml
    row = int(cell[1:])
    m = re.search(r'<((?:\w+:)?row)\b([^>]*?\br="%d"[^>]*?)(/?)>' % row, xml)
    if not m:
        return xml                                # no such row element: skip the note
    prefix = m.group(1)[:-3]                      # "" or "x:"
    c = (f'<{prefix}c r="{cell}" t="inlineStr"><{prefix}is><{prefix}t xml:space="preserve">'
         f'{escape(NOTE)}</{prefix}t></{prefix}is></{prefix}c>')
    if m.group(3):                                # self-closing <row .../>
        return xml[:m.start()] + f'<{m.group(1)}{m.group(2)}>{c}</{m.group(1)}>' + xml[m.end():]
    return xml[:m.end()] + c + xml[m.end():]

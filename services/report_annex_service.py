"""
report_annex_service.py

Builds the narrative content for the workbook's two annex sheets — SWOT and
Conclusion — as ordinary "Sheet!Cell" -> text answers, so the existing template-fill
mechanism writes them with everything else (no new write path).

The SWOT quadrants come from the same swot_agent that feeds the Word report; here we
parse its four sections into the four quadrant cells. The Conclusion narrative is
synthesised from the model's own figures (revenue trajectory, DSCR, profitability),
so it never quotes a number the Excel disagrees with and needs no extra LLM call.

These cells target the bank_loan CMA workbook's SWOT / Conclusion sheets. For any
template that does not have those sheets the keys simply find no sheet and are
skipped by fill_template — safe to always include.
"""

from __future__ import annotations

import logging
import re

logger = logging.getLogger("report_annex")

# Cell map — must match the sheets added to CMA_Dashboard_Premium.xlsx.
SWOT_CELLS = {
    "strengths": "SWOT!B6",
    "weaknesses": "SWOT!D6",
    "opportunities": "SWOT!B8",
    "threats": "SWOT!D8",
}
CONCLUSION_CELL = "Conclusion!B28"

_SECTIONS = ("strengths", "weaknesses", "opportunities", "threats")


def _parse_swot(md: str) -> dict:
    """Split swot_agent markdown into four bullet blocks. Tolerant of '### Strengths',
    '**Strengths**', 'Strengths:' and of '-', '*' or numbered bullets."""
    out = {k: [] for k in _SECTIONS}
    if not md:
        return out
    current = None
    for raw in md.splitlines():
        line = raw.strip()
        low = re.sub(r"[^a-z]", "", line.lower())
        matched = None
        for sec in _SECTIONS:
            # a heading line that is essentially just the section word
            if low.startswith(sec) and len(low) <= len(sec) + 2:
                matched = sec
                break
        if matched:
            current = matched
            continue
        if current and line:
            item = re.sub(r"^[-*•\d.)\s]+", "", line).strip()
            # skip markdown table rows / separators
            if item and not item.startswith("|") and not set(item) <= set("-|: "):
                out[current].append(item)
    return out


def _bullets(items: list, limit: int = 6) -> str:
    """Quadrant text: up to `limit` concise bullet lines."""
    picked = [i for i in items if i][:limit]
    return "\n".join(f"•  {i}" for i in picked)


def _fmt_inr(v) -> str:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return str(v)
    a = abs(v)
    if a >= 1e7:
        return f"₹{v/1e7:,.2f} Cr"
    if a >= 1e5:
        return f"₹{v/1e5:,.2f} L"
    return f"₹{v:,.0f}"


def verdict_tier(kpis: dict) -> tuple[str, str] | None:
    """The same tier + one-line reason the workbook's own Conclusion!D25 formula
    computes, from the same kpis dict (see _real_kpis_from_recalc) -- the ONE place
    this is decided, so the Excel verdict, the Word conclusion paragraph and the
    narrative-writing prompt can never independently reach three different answers
    (5-report audit bug B: the narrative twice quoted a DSCR the workbook did not show,
    and called a REVIEW-REQUIRED report "financially viable"). Returns None if there
    isn't enough data to judge (no DSCR available)."""
    if not isinstance(kpis, dict):
        return None
    if kpis.get("scale_flag"):
        return ("REVIEW REQUIRED",
                "the revenue/cost scale, EBITDA margin or DSCR falls outside the normal "
                "range for a proposal this size and needs manual verification")
    avg_dscr = kpis.get("avg_dscr")
    if avg_dscr is None:
        return None
    try:
        d = float(avg_dscr)
        m = float(kpis.get("min_dscr")) if kpis.get("min_dscr") is not None else d
    except (TypeError, ValueError):
        return None
    if d >= 1.5 and m >= 1.25:
        return ("STRONG", f"average DSCR {d:.2f} and minimum-year DSCR {m:.2f} both clear "
                          f"the bank's comfort threshold")
    if d >= 1.2 and m >= 1.0:
        return ("BANKABLE", f"average DSCR {d:.2f} meets the 1.20 minimum with no year "
                            f"below 1.00 (weakest year {m:.2f})")
    return ("BELOW NORM", f"average DSCR {d:.2f} or the weakest year ({m:.2f}) falls "
                          f"short of the bank's floor")


def _conclusion_text(project, kpis: dict) -> str:
    """A professional, data-grounded conclusion paragraph built from the model's own
    figures. kpis: optional {'revenue_y1','revenue_y5','pat_y5','avg_dscr','min_dscr',
    'scale_flag',...} — see _real_kpis_from_recalc, the normal source for these."""
    name = getattr(project, "title", None) or "The project"
    industry = getattr(project, "industry", None) or "the sector"
    parts = []
    parts.append(
        f"{name} has been appraised as a {industry} venture over a five-year projection "
        f"horizon on the basis of the assumptions and financial statements in this workbook."
    )

    rev1, rev5 = kpis.get("revenue_y1"), kpis.get("revenue_y5")
    if rev1 and rev5:
        try:
            growth = (float(rev5) / float(rev1) - 1) * 100 if float(rev1) else 0
            parts.append(
                f"Revenue is projected to grow from {_fmt_inr(rev1)} in Year 1 to "
                f"{_fmt_inr(rev5)} by Year 5 ({growth:+.0f}% over the period), reflecting the "
                f"capacity build-up and demand assumptions adopted."
            )
        except (TypeError, ValueError):
            pass

    # verdict_tier runs the SAME test the workbook's own Conclusion!D25 formula runs, on
    # the same recalculated figures — so this paragraph and the sheet can never disagree,
    # unlike before when this text had no working access to the real numbers at all (see
    # _real_kpis_from_recalc) and the sheet was the only place a verdict was computed.
    tier = verdict_tier(kpis)
    scale_flag = kpis.get("scale_flag")
    avg_dscr, min_dscr = kpis.get("avg_dscr"), kpis.get("min_dscr")
    if scale_flag:
        parts.append(
            "The projected revenue, margin or coverage figures fall outside the normal "
            "range for a proposal of this size and warrant manual verification before "
            "these numbers are relied on — REVIEW REQUIRED."
        )
    elif tier and avg_dscr is not None:
        try:
            d = float(avg_dscr)
            m = float(min_dscr) if min_dscr is not None else d
            if tier[0] == "STRONG":
                verdict = (f"comfortably serviceable — the average DSCR ({d:.2f}) is well "
                           f"above the 1.20 benchmark and the weakest year ({m:.2f}) still "
                           f"clears 1.25")
            elif tier[0] == "BANKABLE":
                verdict = (f"serviceable — the average DSCR ({d:.2f}) meets the 1.20 minimum "
                           f"and no year falls below 1.00 (weakest year {m:.2f})")
            else:
                verdict = (f"below the benchmark banks look for (average DSCR {d:.2f}, "
                           f"weakest year {m:.2f}), indicating the debt structure or margins "
                           f"should be revisited before sanction")
            parts.append(f"The debt is {verdict}.")
        except (TypeError, ValueError):
            pass

    pat5 = kpis.get("pat_y5")
    if pat5 is not None:
        try:
            p = float(pat5)
            if p > 0:
                parts.append(
                    f"The project turns a Year-5 profit after tax of {_fmt_inr(p)}, and the "
                    f"cash accruals support the projected repayment schedule."
                )
            else:
                parts.append(
                    "The project does not reach a positive Year-5 profit after tax on the "
                    "current assumptions; pricing, cost or scale assumptions warrant review."
                )
        except (TypeError, ValueError):
            pass

    # The closing line asserts viability only when tier says so -- it used to default to
    # "financially viable" whenever scale_flag was falsy, which is also what an EMPTY kpis
    # dict gives (no data at all, e.g. the recalc did not run) -- so a report with a
    # REVIEW REQUIRED verdict, or simply no verdict computed yet, still closed on a
    # confident "financially viable" (5-report audit follow-up: the Excel Conclusion
    # sheet's own BANKABILITY VERDICT went dynamic in bug 4, but this paragraph — which
    # both the sheet and the Word report show — never became conditional on it).
    if scale_flag:
        parts.append(
            "On this basis the proposal cannot yet be confirmed as viable: the figures "
            "above should be checked against the promoter's actual capacity, pricing and "
            "cost inputs before this report is relied upon for a sanction decision."
        )
    elif tier and tier[0] in ("STRONG", "BANKABLE"):
        parts.append(
            "On the strength of the projected financials, ratios and coverage set out above, "
            "the proposal is considered financially viable, subject to the assumptions holding "
            "and the usual terms of sanction."
        )
    elif tier:  # BELOW NORM
        parts.append(
            "On the figures above, the proposal's coverage falls short of the bank's norm; "
            "the debt structure, margins or promoter contribution should be revisited before "
            "a viability conclusion can be drawn."
        )
    else:
        # No DSCR/verdict data was available to this paragraph at all -- say so plainly
        # rather than asserting a conclusion the figures were never checked against.
        parts.append(
            "The proposal's viability rests on the projected financials, ratios and coverage "
            "in this workbook; refer to the DSCR schedule and the Bankability Verdict above "
            "for the specific conclusion those figures support."
        )
    return "  ".join(parts)


def _real_kpis_from_recalc(recalc_bytes: bytes) -> dict:
    """The same figures, read from the same cells, that the workbook's own
    Conclusion!D25 verdict formula uses — so the Word conclusion and the Excel verdict
    are always computed from one source. Returns {} if recalc_bytes is unavailable or
    the expected sheets/cells are not present (never raises)."""
    if not recalc_bytes:
        return {}
    try:
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(recalc_bytes), data_only=True)
        if "DSCR" not in wb.sheetnames or "Annual_Summary" not in wb.sheetnames:
            return {}
        dscr_row = [wb["DSCR"][f"{c}14"].value for c in "CDEFG"]
        dscr_row = [float(v) for v in dscr_row if isinstance(v, (int, float))]
        if not dscr_row:
            return {}
        avg_dscr = sum(dscr_row) / len(dscr_row)
        min_dscr = min(dscr_row)

        annual = wb["Annual_Summary"]
        sales = [annual[f"{c}8"].value for c in "CDEFG"]
        ebitda = [annual[f"{c}22"].value for c in "CDEFG"]
        pat = [annual[f"{c}26"].value for c in "CDEFG"]
        rev1 = sales[0] if isinstance(sales[0], (int, float)) else None
        rev5 = sales[4] if isinstance(sales[4], (int, float)) else None
        pat5 = pat[4] if isinstance(pat[4], (int, float)) else None

        margins = [e / s for e, s in zip(ebitda, sales)
                   if isinstance(e, (int, float)) and isinstance(s, (int, float)) and s]
        max_margin = max(margins) if margins else None

        proj_cost = None
        if "Assumptions" in wb.sheetnames:
            loan = wb["Assumptions"]["C8"].value
            equity = wb["Assumptions"]["C9"].value
            if isinstance(loan, (int, float)) and isinstance(equity, (int, float)):
                proj_cost = loan + equity
        ratio = (rev1 / proj_cost) if (rev1 and proj_cost) else None

        # Same three flags as Conclusion!D25 (bug 1 from the 5-report audit: a scale/unit
        # mismatch inflates sales while DSCR/margin quietly go non-sensical) — checked
        # here too so the Word narrative never asserts "STRONG"/"viable" over figures the
        # workbook itself would mark REVIEW REQUIRED.
        scale_flag = bool(
            (max_margin is not None and max_margin > 0.4)
            or (ratio is not None and (ratio < 0.3 or ratio > 8))
            or (avg_dscr > 10)
        )
        return {"revenue_y1": rev1, "revenue_y5": rev5, "pat_y5": pat5,
                "avg_dscr": avg_dscr, "min_dscr": min_dscr, "scale_flag": scale_flag}
    except Exception:
        logger.warning("annex: could not read real KPIs from the recalculated workbook",
                       exc_info=True)
        return {}


def _kpis_from_model(model: dict) -> dict:
    """Pull the handful of figures the conclusion needs out of the stored model, if
    present. Best-effort — any missing field simply drops its sentence."""
    k = {}
    if not isinstance(model, dict):
        return k
    fm = model.get("financials") or model.get("engine") or model
    # common shapes: annual lists under 'revenue'/'pat', or a ratios block
    def _first_last(seq):
        if isinstance(seq, (list, tuple)) and seq:
            return seq[0], seq[-1]
        return None, None
    rev = (fm.get("revenue") or fm.get("annual_revenue")
           or (fm.get("profit") or {}).get("revenue"))
    r1, r5 = _first_last(rev)
    k["revenue_y1"], k["revenue_y5"] = r1, r5
    pat = (fm.get("pat") or (fm.get("profit") or {}).get("pat"))
    _, p5 = _first_last(pat)
    k["pat_y5"] = p5
    ratios = fm.get("ratios") or {}
    k["avg_dscr"] = ratios.get("average_dscr") or model.get("avg_dscr")
    return {kk: vv for kk, vv in k.items() if vv is not None}


def build_annex_cell_answers(project, purpose_label: str = "", model: dict = None,
                             swot_markdown: str = None, recalc_bytes: bytes = None) -> dict:
    """Return {"SWOT!B6": ..., ..., "Conclusion!B28": ...} for the workbook annex.

    swot_markdown: pass the swot_agent output if already computed; else it is
    generated here. Failures are non-fatal — a missing quadrant just stays blank.

    recalc_bytes: the server-recalculated workbook, when available. The narrative
    `model` dict has not carried real financial figures since the "sheets" array was
    dropped from the LLM prompt (see financial_model_service.py) — so without this the
    conclusion's DSCR/revenue sentences silently never fired, on every report, for
    however long that has been true. Reading the real numbers back out of the same
    recalculated workbook the Excel verdict formula reads is also what lets this
    paragraph and Conclusion!D25 agree instead of being two independent guesses."""
    answers = {}

    # SWOT
    try:
        md = swot_markdown
        if md is None:
            from agents.swot_agent import swot_agent
            md = swot_agent(
                business_name=getattr(project, "title", "") or "",
                industry=getattr(project, "industry", "") or "",
                country=getattr(project, "country", "") or "",
                description=getattr(project, "project_description", "") or "",
            )
        quad = _parse_swot(md)
        for sec, cell in SWOT_CELLS.items():
            text = _bullets(quad.get(sec, []))
            if text:
                answers[cell] = text
    except Exception:
        logger.warning("annex: SWOT build failed", exc_info=True)

    # Conclusion
    try:
        kpis = _real_kpis_from_recalc(recalc_bytes) or _kpis_from_model(model or {})
        answers[CONCLUSION_CELL] = _conclusion_text(project, kpis)
    except Exception:
        logger.warning("annex: conclusion build failed", exc_info=True)

    return answers

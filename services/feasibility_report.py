"""
feasibility_report.py

Reads the recalculated Feasibility Study workbook into the pieces of the stored model the
Word report and the short report are written from: financial_summary, statement_tables,
key_assumptions, and the real KPIs behind the verdict. Feasibility only --
generation_router calls this for the feasibility_study_all_industries_* template.

Why a reader of its own: the generic readers (financial_summary_service,
report_annex_service._real_kpis_from_recalc) look for the CMA workbook's sheets and take
every figure as rupees. This workbook keeps money in Rs LAKH on sheets of its own, so they
either find nothing or report a project a hundred thousand times too small. Here every
figure is read from a fixed cell and converted to rupees, and the rows are labelled the
way the report builders already recognise (the CMA vocabulary), so neither builder needs
to change.
"""

from __future__ import annotations

import logging
from io import BytesIO

from openpyxl import load_workbook

from services import feasibility_template_io as fio

logger = logging.getLogger("feasibility")

LAKH = 100000
YEARS = 5
YEAR_COLS = ["C", "D", "E", "F", "G"]          # Year 1..5 on the annual sheets
BS_COLS = ["D", "E", "F", "G", "H"]            # Balance Sheet / Cash Flow: C is "At start"

PNL = "Profitability (Annual)"
RATIOS = "Ratios & DSCR"


def _open(recalc_bytes):
    try:
        wb = load_workbook(BytesIO(recalc_bytes), data_only=True)
    except Exception:
        logger.warning("feasibility_report: could not open the recalculated workbook",
                       exc_info=True)
        return None
    return wb if PNL in wb.sheetnames and RATIOS in wb.sheetnames else None


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _row(ws, row, cols=YEAR_COLS, scale=1.0):
    out = []
    for c in cols:
        n = _num(ws[f"{c}{row}"].value)
        out.append(None if n is None else n * scale)
    return out


def _rupees(ws, row, cols=YEAR_COLS):
    return _row(ws, row, cols, LAKH)


def _text(ws, cell, default=""):
    v = ws[cell].value
    return v.strip() if isinstance(v, str) and v.strip() else default


def _first(series):
    return next((v for v in series if v is not None), None)


def _last(series):
    return next((v for v in reversed(series) if v is not None), None)


def _avg(series):
    vals = [v for v in series if v is not None]
    return sum(vals) / len(vals) if vals else None


def _input(wb, name):
    sheet, cell, _ = fio.INPUTS[name]
    return _num(wb[sheet][cell].value)


def _has_loan(wb) -> bool:
    return (_input(wb, "term_loan") or 0) > 0


# ── financial_summary ───────────────────────────────────────────────────────────
def financial_summary(recalc_bytes: bytes) -> dict:
    """Same shape as financial_summary_service.extract_financial_summary, in rupees."""
    wb = _open(recalc_bytes)
    if wb is None:
        return {}
    p, r = wb[PNL], wb[RATIOS]
    series = {
        "Net Sales / Revenue": _rupees(p, 9),
        "EBITDA": _rupees(p, 27),
        "EBIT (Operating Profit)": _rupees(p, 30),
        "Profit Before Tax": _rupees(p, 34),
        "Profit After Tax": _rupees(p, 36),
        "Cash Accrual": _rupees(p, 38),
    }
    ratios = {
        "Current Ratio": _row(r, 16),
        "Debt–Equity Ratio": _row(r, 12),
        "EBITDA Margin": _row(p, 28),
        "Net Profit Margin": _row(p, 37),
    }
    # No term loan, no debt service: a DSCR row of zeros reads as broken, not as
    # "not applicable", so it is left out (as the generic reader does).
    if _has_loan(wb):
        ratios = {"DSCR": _row(r, 25), **ratios}
    series = {k: v for k, v in series.items() if any(x is not None for x in v)}
    ratios = {k: v for k, v in ratios.items() if any(x is not None for x in v)}
    if not series:
        return {}
    revenue = series.get("Net Sales / Revenue", [None] * YEARS)
    return {
        "years": [f"Year {i + 1}" for i in range(YEARS)],
        "series": series,
        "ratios": ratios,
        "cards": {
            "revenue_y1": _first(revenue),
            "revenue_y5": _last(revenue),
            "ebitda_y5": _last(series.get("EBITDA", [None])),
            "pat_y5": _last(series.get("Profit After Tax", [None])),
            "avg_dscr": _avg(ratios.get("DSCR", [None])),
            "net_margin_y5": _last(ratios.get("Net Profit Margin", [None])),
        },
    }


# ── real KPIs behind the verdict ────────────────────────────────────────────────
def real_kpis(recalc_bytes: bytes) -> dict:
    """Same shape as report_annex_service._real_kpis_from_recalc, so verdict_tier() gives
    the same STRONG / BANKABLE / BELOW NORM / REVIEW REQUIRED tiers on the same rules."""
    wb = _open(recalc_bytes)
    if wb is None:
        return {}
    p, r = wb[PNL], wb[RATIOS]
    sales, ebitda, pat = _row(p, 9), _row(p, 27), _row(p, 36)
    rev1, rev5, pat5 = sales[0], sales[4], pat[4]
    if rev1 is None:
        return {}
    out = {"revenue_y1": rev1 * LAKH,
           "revenue_y5": rev5 * LAKH if rev5 is not None else None,
           "pat_y5": pat5 * LAKH if pat5 is not None else None}

    margins = [e / s for e, s in zip(ebitda, sales) if e is not None and s]
    max_margin = max(margins) if margins else None
    cost = _num(wb["Project Cost"]["C14"].value)
    ratio = rev1 / cost if cost else None

    avg_dscr = None
    if _has_loan(wb):
        dscr = [v for v in _row(r, 25) if v is not None]
        if dscr:
            avg_dscr = sum(dscr) / len(dscr)
            out["avg_dscr"], out["min_dscr"] = avg_dscr, min(dscr)
    out["scale_flag"] = bool(
        (max_margin is not None and max_margin > 0.4)
        or (ratio is not None and (ratio < 0.3 or ratio > 8))
        or (avg_dscr is not None and avg_dscr > 10)
    )
    return out


def consistency_checks(recalc_bytes: bytes) -> list:
    """The workbook's own tie-outs, in the shape template_analysis.run_checks returns."""
    wb = _open(recalc_bytes)
    if wb is None:
        return []
    out = []

    def add(name, sheet, cell, values):
        vals = [v for v in values if v is not None]
        ok = bool(vals) and all(abs(v) <= 0.01 for v in vals)
        out.append({"name": name, "sheet": sheet, "cell": cell,
                    "value": max((abs(v) for v in vals), default=None),
                    "ok": ok, "tolerance": 0.01})

    add("Project cost = means of finance", "Project Cost", "C21",
        [_num(wb["Project Cost"]["C21"].value)])
    add("Balance sheet tallies (assets = liabilities)", "Balance Sheet", "C21:H21",
        _row(wb["Balance Sheet"], 21, ["C"] + BS_COLS))
    add("Cash flow reconciles to the balance sheet", "Cash Flow", "C21:H21",
        _row(wb["Cash Flow"], 21, ["C"] + BS_COLS))
    return out


# ── statement tables ────────────────────────────────────────────────────────────
def _line(label, vals, kind="money"):
    """One table row. Money comes out of the workbook in lakh and goes into the report in
    rupees. A %-row the report would print as 0.55 is shown as 55.0 with "(%)" on the
    label, unless the label says "margin" (the Word builder already shows those as %)."""
    if kind == "money":
        vals = [None if v is None else v * LAKH for v in vals]
    elif kind == "pct" and "margin" not in label.lower():
        vals = [None if v is None else round(v * 100, 2) for v in vals]
        label = f"{label} (%)"
    return (label, vals, False)


def _heading(label):
    return (label, [None] * YEARS, True)


def _labelled(ws, row, default):
    """Row labels that the workbook builds from the industry (Industry Labels) come back
    as text after recalculation; fall back to a plain name if one did not."""
    return _text(ws, f"A{row}", default)


def statement_tables(recalc_bytes: bytes) -> list:
    """[{key, title, subtitle, rows}] -- money in rupees, rows labelled the way the short
    report and the Word builder look for them."""
    wb = _open(recalc_bytes)
    if wb is None:
        return []
    p, r = wb[PNL], wb[RATIOS]
    wc, bs, cf, be = (wb["Working Capital"], wb["Balance Sheet"], wb["Cash Flow"],
                      wb["Break-even & Sensitivity"])
    unit = _text(p, "B7", "units")
    tables = []

    # The consolidated summary the short report reads (by these headings) for volume,
    # the cost lines and the profit rows.
    summary = [
        _heading("PRODUCTION & SALES"),
        (f"Output / sales volume — {unit} (units)", _row(p, 7), False),
        _line("Average capacity utilisation", _row(p, 6), "pct"),
        _line("Net Sales (net of GST)", _row(p, 9)),
        _heading("COST STRUCTURE"),
        _line("Raw material / direct cost", _row(p, 11)),
        _line("Power & utilities (per unit)", _row(p, 12)),
        _line("Consumables & packing", _row(p, 13)),
        _line("Direct labour", _row(p, 14)),
        _line("Admin & management salaries", _row(p, 19)),
        _line("Fixed utilities & establishment", _row(p, 20)),
        _line("Repairs & maintenance", _row(p, 21)),
        _line("Insurance", _row(p, 22)),
        _line("Other admin overheads", _row(p, 23)),
        _line("Selling & distribution", _row(p, 24)),
        _heading("PROFITABILITY"),
        _line("EBITDA", _row(p, 27)),
        _line("EBITDA margin", _row(p, 28), "pct"),
        _line("Depreciation & amortisation", _row(p, 29)),
        _line("Interest & finance cost", _row(p, 33)),
        _line("Profit before tax (PBT)", _row(p, 34)),
        _line("Provision for tax", _row(p, 35)),
        _line("Profit After Tax (PAT)", _row(p, 36)),
        _line("Cash Accrual (PAT + depreciation)", _row(p, 38)),
    ]
    tables.append({"key": "feas_summary", "title": "Consolidated Five-Year Summary",
                   "subtitle": "Volume, sales, cost structure and profitability by year",
                   "rows": summary})

    operating = [
        _line("Net Sales (net of GST)", _row(p, 9)),
        _line(_labelled(p, 11, "Direct cost"), _row(p, 11)),
        _line(_labelled(p, 12, "Utilities (variable)"), _row(p, 12)),
        _line(_labelled(p, 13, "Consumables & packing"), _row(p, 13)),
        _line("Direct labour", _row(p, 14)),
        _line("Total variable cost", _row(p, 15)),
        _line("Contribution", _row(p, 16)),
        _line("Contribution margin", _row(p, 17), "pct"),
        _line("Total fixed overheads", _row(p, 25)),
        _line("EBITDA (Operating profit)", _row(p, 27)),
        _line("Depreciation & amortisation", _row(p, 29)),
        _line("Interest on term loan", _row(p, 31)),
        _line("Interest on working capital", _row(p, 32)),
        _line("Profit before tax (PBT)", _row(p, 34)),
        _line("Provision for tax", _row(p, 35)),
        _line("Profit After Tax (PAT)", _row(p, 36)),
        _line("PAT margin", _row(p, 37), "pct"),
    ]
    tables.append({"key": "feas_operating", "title": "Projected Profitability",
                   "subtitle": "Operating statement for the five projected years",
                   "rows": operating})

    if _has_loan(wb):
        repay = [
            _line("Interest on term loan", _row(p, 31)),
            _line("Principal repayment", _row(cf, 15, BS_COLS)),
            _line("Closing balance", _row(bs, 9, BS_COLS)),
        ]
        tables.append({"key": "feas_repayment", "title": "Term Loan Repayment Schedule",
                       "subtitle": "Interest, principal and the balance outstanding",
                       "rows": repay})

    balance = [
        _heading("LIABILITIES / SOURCES OF FUNDS"),
        _line("Promoter's capital / equity", _row(bs, 6, BS_COLS)),
        _line("Reserves & surplus", _row(bs, 7, BS_COLS)),
        _line("Net worth", _row(bs, 8, BS_COLS)),
        _line("Term loan (closing)", _row(bs, 9, BS_COLS)),
        _line("Working-capital borrowing", _row(bs, 10, BS_COLS)),
        _line("Sundry creditors", _row(bs, 11, BS_COLS)),
        _line("Total liabilities", _row(bs, 12, BS_COLS)),
        _heading("ASSETS / APPLICATION OF FUNDS"),
        _line("Net fixed assets (incl. land)", _row(bs, 14, BS_COLS)),
        _line("Pre-operative expenses (unamortised)", _row(bs, 15, BS_COLS)),
        _line(_labelled(bs, 16, "Stock (raw material)"), _row(bs, 16, BS_COLS)),
        _line(_labelled(bs, 17, "Stock (finished goods)"), _row(bs, 17, BS_COLS)),
        _line("Receivables (debtors)", _row(bs, 18, BS_COLS)),
        _line("Cash & bank balance", _row(bs, 19, BS_COLS)),
        _line("Total assets", _row(bs, 20, BS_COLS)),
    ]
    tables.append({"key": "feas_balance_sheet", "title": "Projected Balance Sheet",
                   "subtitle": "Position at the close of each year", "rows": balance})

    cash = [
        _line("Cash accruals (PAT + depreciation)", _row(cf, 6, BS_COLS)),
        _line("Increase in cash-credit borrowing", _row(cf, 9, BS_COLS)),
        _line("Increase in creditors", _row(cf, 10, BS_COLS)),
        _line("Total sources", _row(cf, 11, BS_COLS)),
        _line("Increase in current assets", _row(cf, 14, BS_COLS)),
        _line("Term-loan principal repayment", _row(cf, 15, BS_COLS)),
        _line("Total uses", _row(cf, 16, BS_COLS)),
        _line("Net surplus / (deficit) for the year", _row(cf, 18, BS_COLS)),
        _line("Closing cash balance", _row(cf, 20, BS_COLS)),
    ]
    tables.append({"key": "feas_cash_flow", "title": "Projected Cash Flow",
                   "subtitle": "Sources and uses of cash, year by year", "rows": cash})

    working = [
        _line("Total current assets", _row(wc, 18)),
        _line(_labelled(wc, 20, "Sundry creditors"), _row(wc, 20)),
        _line("Working-capital gap", _row(wc, 22)),
        _line("Margin from promoter (25% of current assets)", _row(wc, 23)),
        _line("Assessed bank finance (cash-credit limit)", _row(wc, 24)),
        _line("Interest on cash credit", _row(wc, 27)),
    ]
    tables.append({"key": "feas_working_capital", "title": "Working Capital Assessment",
                   "subtitle": "Current assets, the gap and the bank finance it supports",
                   "rows": working})

    ratios = [
        _line("EBITDA margin", _row(r, 6), "pct"),
        _line("PAT margin", _row(r, 8), "pct"),
        _line("Return on capital employed", _row(r, 9), "pct"),
        _line("Return on net worth", _row(r, 10), "pct"),
        ("Debt : Equity (x)", _row(r, 12), False),
        ("TOL / TNW (x)", _row(r, 13), False),
        ("Interest coverage (x)", _row(r, 14), False),
        ("Current ratio (x)", _row(r, 16), False),
        ("Debtors (days)", _row(r, 20), False),
    ]
    if _has_loan(wb):
        ratios += [
            _line("Cash available for debt service", _row(r, 23)),
            _line("Debt obligation (principal + interest)", _row(r, 24)),
            ("DSCR (x)", _row(r, 25), False),
        ]
    tables.append({"key": "feas_ratios", "title": "Financial Ratios and Debt-Service Coverage",
                   "subtitle": "Profitability, leverage, liquidity and repayment capacity",
                   "rows": ratios})

    breakeven = [
        _line("Net sales", _row(be, 6)),
        _line("Contribution", _row(be, 7)),
        _line("Fixed cost (overheads, depreciation, interest)", _row(be, 9)),
        _line("Break-even sales", _row(be, 10)),
        _line("Break-even as share of sales", _row(be, 11), "pct"),
        _line("Margin of safety", _row(be, 12), "pct"),
        _line("Cash break-even sales", _row(be, 13)),
    ]
    tables.append({"key": "feas_breakeven", "title": "Break-even Analysis",
                   "subtitle": "The sales level at which the project stops losing money",
                   "rows": breakeven})

    # A table needs at least two real figures to be worth a page.
    return [t for t in tables
            if sum(1 for _, v, h in t["rows"] if not h and any(x is not None for x in v)) >= 2]


# ── key assumptions ─────────────────────────────────────────────────────────────
def _money(v):
    return None if v is None else f"{v * LAKH:,.0f}"


def _pct(v):
    return None if v is None else f"{v * 100:.2f}%"


def _plain(v, places=2):
    if v is None:
        return None
    return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.{places}f}"


# pc_* input -> how the cost line is named. "— cost" is what the short report reads as a
# capital line; every other label below avoids the words it treats as capital.
_COST_LINES = [
    ("pc_land", "Land — cost"),
    ("pc_building", "Building & civil works — cost"),
    ("pc_equipment", "Plant, machinery & equipment — cost"),
    ("pc_electrical", "Electrical installation — cost"),
    ("pc_utilities", "Utilities & services — cost"),
    ("pc_misc_assets", "Furniture & other assets — cost"),
    ("pc_contingency", "Contingency — cost"),
    ("pc_preoperative", "Pre-operative expenses — cost"),
    ("pc_wc_margin", "Margin money for working capital — cost"),
]


def key_assumptions(recalc_bytes: bytes) -> list:
    """(label, shown value) pairs, as financial_summary_service.extract_key_assumptions
    returns them: money in rupees, rates as "12.00%". A row with value None is a heading."""
    wb = _open(recalc_bytes)
    if wb is None:
        return []
    g = lambda n: _input(wb, n)                      # noqa: E731
    a = wb["Assumptions"]
    unit = _text(wb[PNL], "B7", "unit")
    util_y1 = [g(f"util_y1_m{m}") for m in range(1, 13)]
    util_y1 = [u for u in util_y1 if u is not None]

    rows = [("A. PROJECT COST & MEANS OF FINANCE", None)]
    for name, label in _COST_LINES:
        v = g(name)
        if v:
            rows.append((label, _money(v)))
    rows += [
        ("TOTAL COST OF PROJECT", _money(_num(wb["Project Cost"]["C14"].value))),
        ("Term Loan applied", _money(g("term_loan"))),
        ("Promoter's Capital / Equity", _money(g("promoter_equity"))),
        ("Interest rate on Term Loan", _pct(g("term_loan_interest"))),
        ("Loan tenure (months)", _plain(g("term_loan_repayment_months"))),
        ("Moratorium (months)", _plain(g("term_loan_moratorium_months"))),
        ("Interest rate on Working Capital", _pct(g("cc_interest"))),
        ("B. CAPACITY, PRICE & COST PER UNIT", None),
        ("Industry", _text(a, "B3") or None),
        (f"Installed capacity ({unit} per month)", _plain(g("capacity_per_month"))),
        ("Working days per month", _plain(g("working_days_per_month"))),
        (f"Selling price (₹ per {unit})", _plain(g("selling_price_per_unit"))),
        (f"Direct cost (₹ per {unit})", _plain(g("direct_cost_per_unit"))),
        (f"Utilities (₹ per {unit})", _plain(g("utilities_per_unit"))),
        (f"Consumables (₹ per {unit})", _plain(g("consumables_per_unit"))),
        (f"Packing (₹ per {unit})", _plain(g("packing_per_unit"))),
        ("Capacity utilisation — Year 1 average",
         _pct(sum(util_y1) / len(util_y1)) if util_y1 else None),
        ("Capacity utilisation — Year 2", _pct(g("util_y2"))),
        ("Capacity utilisation — Year 3", _pct(g("util_y3"))),
        ("Capacity utilisation — Year 4", _pct(g("util_y4"))),
        ("Capacity utilisation — Year 5", _pct(g("util_y5"))),
        ("C. OPERATING COSTS", None),
        ("Direct labour (₹ a month)", _money(g("direct_labour_per_month"))),
        ("Admin & management salaries (₹ a month)", _money(g("admin_salaries_per_month"))),
        ("Fixed utilities (₹ a month)", _money(g("fixed_utility_per_month"))),
        ("Other admin overheads (₹ a month)", _money(g("other_admin_per_month"))),
        ("Repairs & maintenance (% of equipment)", _pct(g("repairs_pct_of_equipment"))),
        ("Insurance (% of assets)", _pct(g("insurance_pct_of_assets"))),
        ("Selling & distribution (% of sales)", _pct(g("selling_dist_pct_of_sales"))),
        ("D. GROWTH, WORKING CAPITAL & TAX", None),
        ("Selling-price escalation a year", _pct(g("price_escalation"))),
        ("Direct-cost escalation a year", _pct(g("direct_cost_escalation"))),
        ("Wage escalation a year", _pct(g("wage_escalation"))),
        ("Raw-material stock (days)", _plain(g("stock_days"))),
        ("Finished-goods stock (days)", _plain(g("finished_stock_days"))),
        ("Credit to customers (days)", _plain(g("debtor_days"))),
        ("Credit from suppliers (days)", _plain(g("creditor_days"))),
        ("Promoter margin on current assets", _pct(g("wc_margin"))),
        ("Income-tax rate", _pct(g("tax_rate"))),
        ("GST rate", _pct(g("gst_rate"))),
    ]
    rows = [(k, v) for k, v in rows if v is not None or k[:2].rstrip(".").isalpha()
            and k.isupper()]
    while rows and rows[-1][1] is None:
        rows.pop()
    return rows

"""
feasibility_template_io.py

Fill the Feasibility Study workbook with inputs and read the results back.

The workbook does the maths (about 3,200 live formulas). This module only:
  1. knows WHICH cells are inputs (INPUTS) and which are results (OUTPUTS),
  2. validates the values before they are written,
  3. writes them, optionally recalculates, and reads the results.

Nothing here touches any other purpose (Bank Loan / CMA etc.). It is not wired into
any router; call it from the Feasibility flow only.

UNITS (same as the workbook): money is in Rs LAKH, except the per-unit prices
(selling price, direct cost, utilities, consumables, packing), which are in Rs per unit.
Every percentage is a FRACTION (0.25 means 25 percent), never 25.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from typing import Dict, Iterable, List, Optional

from openpyxl import load_workbook

TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "templates", "feasibility_study",
    "CMA_Feasibility_Study_All_Industries_Blank_Template.xlsx")

DEFAULT_START_YEAR = 2026

# key -> (sheet, cell, kind).  kind: number | percent | integer | year | text
INPUTS: Dict[str, tuple] = {
    "industry":               ("Assumptions", "B3", "text"),
    "start_year":             ("Assumptions", "B6", "year"),
    "capacity_per_month":     ("Assumptions", "B5", "number"),
    "working_days_per_month": ("Assumptions", "B7", "integer"),
    "selling_price_per_unit": ("Assumptions", "B10", "number"),
    "direct_cost_per_unit":   ("Assumptions", "B11", "number"),
    # Year-1 utilisation, month by month (B16..B27), then Years 2-5 (B28..B31)
    **{f"util_y1_m{m}": ("Assumptions", f"B{15 + m}", "percent") for m in range(1, 13)},
    "util_y2": ("Assumptions", "B28", "percent"),
    "util_y3": ("Assumptions", "B29", "percent"),
    "util_y4": ("Assumptions", "B30", "percent"),
    "util_y5": ("Assumptions", "B31", "percent"),
    "price_escalation":       ("Assumptions", "B33", "percent"),
    "direct_cost_escalation": ("Assumptions", "B34", "percent"),
    "utility_escalation":     ("Assumptions", "B35", "percent"),
    "wage_escalation":        ("Assumptions", "B36", "percent"),
    "overhead_escalation":    ("Assumptions", "B37", "percent"),
    "utilities_per_unit":     ("Assumptions", "B39", "number"),
    "consumables_per_unit":   ("Assumptions", "B40", "number"),
    "packing_per_unit":       ("Assumptions", "B41", "number"),
    "direct_labour_per_month":  ("Assumptions", "B42", "number"),
    "admin_salaries_per_month": ("Assumptions", "B44", "number"),
    "fixed_utility_per_month":  ("Assumptions", "B45", "number"),
    "repairs_pct_of_equipment": ("Assumptions", "B46", "percent"),
    "insurance_pct_of_assets":  ("Assumptions", "B47", "percent"),
    "other_admin_per_month":    ("Assumptions", "B48", "number"),
    "selling_dist_pct_of_sales": ("Assumptions", "B49", "percent"),
    "term_loan":              ("Assumptions", "B51", "number"),
    "promoter_equity":        ("Assumptions", "B52", "number"),
    "term_loan_interest":     ("Assumptions", "B53", "percent"),
    "term_loan_moratorium_months": ("Assumptions", "B54", "integer"),
    "term_loan_repayment_months":  ("Assumptions", "B55", "integer"),
    "cc_interest":            ("Assumptions", "B56", "percent"),
    "stock_days":             ("Assumptions", "B58", "number"),
    "finished_stock_days":    ("Assumptions", "B59", "number"),
    "debtor_days":            ("Assumptions", "B60", "number"),
    "creditor_days":          ("Assumptions", "B61", "number"),
    "wc_margin":              ("Assumptions", "B62", "percent"),
    "dep_building":           ("Assumptions", "B64", "percent"),
    "dep_equipment":          ("Assumptions", "B65", "percent"),
    "dep_electrical":         ("Assumptions", "B66", "percent"),
    "dep_utilities":          ("Assumptions", "B67", "percent"),
    "dep_misc":               ("Assumptions", "B68", "percent"),
    "preop_amortisation":     ("Assumptions", "B69", "percent"),
    "tax_rate":               ("Assumptions", "B70", "percent"),
    "gst_rate":               ("Assumptions", "B71", "percent"),
    # Project cost (Rs lakh)
    "pc_land":          ("Project Cost", "C5", "number"),
    "pc_building":      ("Project Cost", "C6", "number"),
    "pc_equipment":     ("Project Cost", "C7", "number"),
    "pc_electrical":    ("Project Cost", "C8", "number"),
    "pc_utilities":     ("Project Cost", "C9", "number"),
    "pc_misc_assets":   ("Project Cost", "C10", "number"),
    "pc_contingency":   ("Project Cost", "C11", "number"),
    "pc_preoperative":  ("Project Cost", "C12", "number"),
    "pc_wc_margin":     ("Project Cost", "C13", "number"),
}
# start_year has a default in the workbook; every other input must be supplied.
OPTIONAL = {"start_year"}

# name -> (sheet, row, first_col, last_col)  [one row, several years]  or  (sheet, cell)
OUTPUTS = {
    "total_project_cost": ("Project Cost", "C14"),
    "finance_check":      ("Project Cost", "C21"),     # cost minus finance, must be 0
    "net_sales":          ("Profitability (Annual)", 9, "C", "G"),
    "ebitda":             ("Profitability (Annual)", 27, "C", "G"),
    "ebitda_margin":      ("Profitability (Annual)", 28, "C", "G"),
    "pat":                ("Profitability (Annual)", 36, "C", "G"),
    "cash_accruals":      ("Profitability (Annual)", 38, "C", "G"),
    "dscr":               ("Ratios & DSCR", 25, "C", "G"),
    "average_dscr":       ("Ratios & DSCR", "C26"),
    "current_ratio":      ("Ratios & DSCR", 16, "C", "G"),
    "mpbf":               ("Working Capital", 24, "C", "G"),
    "break_even_sales":   ("Break-even & Sensitivity", 10, "C", "G"),
    "margin_of_safety":   ("Break-even & Sensitivity", 12, "C", "G"),
    "balance_sheet_check": ("Balance Sheet", 21, "C", "H"),   # must be 0 (at start + Yr1-5)
    "cash_flow_check":     ("Cash Flow", 21, "C", "H"),       # must be 0
}


def industries(path: str = TEMPLATE_PATH) -> List[str]:
    """Industry names the workbook accepts (header row of 'Industry Labels')."""
    ws = load_workbook(path, read_only=True)["Industry Labels"]
    row = next(ws.iter_rows(min_row=4, max_row=4))
    return [c.value for c in row[2:] if c.value]


def missing_inputs(values: dict) -> List[str]:
    return [k for k in INPUTS if k not in values and k not in OPTIONAL]


def validate(values: dict, path: str = TEMPLATE_PATH, allow_missing: bool = False) -> None:
    unknown = [k for k in values if k not in INPUTS]
    if unknown:
        raise ValueError(f"unknown input keys: {unknown}")
    if not allow_missing:
        miss = missing_inputs(values)
        if miss:
            raise ValueError(f"missing inputs: {miss}")
    for k, v in values.items():
        kind = INPUTS[k][2]
        if kind == "text":
            if v not in industries(path):
                raise ValueError(f"{k}={v!r} is not one of the workbook's industries")
            continue
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise ValueError(f"{k} must be a number, got {v!r}")
        if v != v or v in (float("inf"), float("-inf")):
            raise ValueError(f"{k} is not a finite number")
        if v < 0:
            raise ValueError(f"{k} cannot be negative ({v})")
        if kind == "percent" and v > 1:
            raise ValueError(f"{k}={v} looks like a whole percent; give a fraction (0.25 for 25%)")
        if kind in ("integer",) and int(v) != v:
            raise ValueError(f"{k} must be a whole number, got {v}")
        if kind == "year" and not (2000 <= v <= 2100):
            raise ValueError(f"{k}={v} is not a plausible calendar year")


def fill_inputs(values: dict, out_path: str, template_path: str = TEMPLATE_PATH,
                allow_missing: bool = False) -> str:
    """Copy the template to out_path and write the inputs. Formulas are untouched."""
    validate(values, template_path, allow_missing)
    wb = load_workbook(template_path)
    for k, v in values.items():
        sheet, cell, _ = INPUTS[k]
        wb[sheet][cell].value = v
    wb.save(out_path)
    return out_path


def recalculate(path: str) -> str:
    """Recalculate with LibreOffice (headless) and overwrite `path` with cached values.
    Raises RuntimeError if LibreOffice is not installed (Excel recalculates on open)."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice not found; open the file in Excel to recalculate")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", tmp, path],
                       check=True, capture_output=True, timeout=180)
        shutil.copyfile(os.path.join(tmp, os.path.basename(path)), path)
    return path


def read_outputs(path: str, names: Optional[Iterable[str]] = None) -> dict:
    """Read results (needs a recalculated file). Single cell -> value; years -> list."""
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in (names or OUTPUTS):
        spec = OUTPUTS[name]
        ws = wb[spec[0]]
        if len(spec) == 2:
            out[name] = ws[spec[1]].value
        else:
            _, row, c1, c2 = spec
            out[name] = [c.value for c in ws[f"{c1}{row}:{c2}{row}"][0]]
    return out

"""The Feasibility Study report path, end to end, without LibreOffice and without AI.

LibreOffice is replaced by `fake_recalc`: it takes the really filled workbook and writes
known figures (Rs lakh, as the workbook keeps them) into the cells the formulas would have
computed. So the input cells are the genuinely filled ones, and every assertion below is
about how those figures reach the report -- in rupees, under the labels the report builders
read. The formulas themselves are not exercised here (that needs LibreOffice).

Every model call is a tripwire: services.claude_service._create raises and is counted, and
the HTTP proxy points at a closed port, so nothing can reach a model or the network.
"""
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

os.environ["DATABASE_URL"] = "sqlite://"
for _k in ("DEEPSEEK_API_KEY", "HEAVY_API_KEY", "OPENAI_API_KEY"):
    os.environ[_k] = "test-dummy"
for _k in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy"):
    os.environ[_k] = "http://127.0.0.1:9"

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)

from openpyxl import load_workbook  # noqa: E402

from services import feasibility_template_io as fio  # noqa: E402
from services import feasibility_report as frep  # noqa: E402
from services.report_annex_service import verdict_tier  # noqa: E402

HAVE_TEMPLATE = os.path.isfile(fio.TEMPLATE_PATH)
FEAS_ID = "feasibility_study_all_industries_6648eea9"
LAKH = 100000

INPUTS = {
    "industry": "Manufacturing", "start_year": 2026, "working_days_per_month": 26,
    "capacity_per_month": 10000, "selling_price_per_unit": 450, "direct_cost_per_unit": 160,
    **{f"util_y1_m{m}": round(0.40 + 0.02 * (m - 1), 2) for m in range(1, 13)},
    "util_y2": 0.70, "util_y3": 0.75, "util_y4": 0.80, "util_y5": 0.80,
    "price_escalation": 0.05, "direct_cost_escalation": 0.05, "utility_escalation": 0.05,
    "wage_escalation": 0.06, "overhead_escalation": 0.05,
    "utilities_per_unit": 15, "consumables_per_unit": 10, "packing_per_unit": 8,
    "direct_labour_per_month": 3.0, "admin_salaries_per_month": 1.5,
    "fixed_utility_per_month": 0.5, "repairs_pct_of_equipment": 0.02,
    "insurance_pct_of_assets": 0.005, "other_admin_per_month": 0.5,
    "selling_dist_pct_of_sales": 0.03, "term_loan": 55, "promoter_equity": 25,
    "term_loan_interest": 0.105, "term_loan_moratorium_months": 6,
    "term_loan_repayment_months": 60, "cc_interest": 0.11, "stock_days": 30,
    "finished_stock_days": 15, "debtor_days": 45, "creditor_days": 30, "wc_margin": 0.25,
    "dep_building": 0.10, "dep_equipment": 0.15, "dep_electrical": 0.10,
    "dep_utilities": 0.15, "dep_misc": 0.10, "preop_amortisation": 0.20,
    "tax_rate": 0.25, "gst_rate": 0.18,
    "pc_land": 0, "pc_building": 20, "pc_equipment": 40, "pc_electrical": 5,
    "pc_utilities": 4, "pc_misc_assets": 3, "pc_contingency": 3, "pc_preoperative": 2,
    "pc_wc_margin": 3,
}

PNL = "Profitability (Annual)"
SALES = [275.4, 396.9, 446.5, 500.1, 525.1]
EBITDA = [40.0, 70.0, 85.0, 100.0, 108.0]
PAT = [20.0, 40.0, 52.0, 65.0, 72.0]
DSCR = [1.40, 1.90, 2.20, 2.50, 2.90]
TL_INTEREST = [5.5, 5.0, 4.0, 2.8, 1.6]
PRINCIPAL = [5.0, 11.0, 11.0, 11.0, 11.0]
TL_CLOSING = [50.0, 39.0, 28.0, 17.0, 6.0]


def cells(values):
    return {f"{fio.INPUTS[k][0]}!{fio.INPUTS[k][1]}": v for k, v in values.items()}


def _years(ws, row, vals, cols=("C", "D", "E", "F", "G")):
    for c, v in zip(cols, vals):
        ws[f"{c}{row}"] = v


def fake_recalc(filled: bytes, sales=SALES, ebitda=EBITDA, dscr=DSCR) -> bytes:
    """What LibreOffice would hand back: the filled workbook with computed cells valued."""
    wb = load_workbook(io.BytesIO(filled))
    p = wb[PNL]
    p["B7"] = "Units"
    p["A11"], p["A12"], p["A13"] = "Raw material consumed", "Power & fuel", "Consumables"
    p["A20"] = "Fixed utilities"
    _years(p, 6, [0.51, 0.70, 0.75, 0.80, 0.80])
    _years(p, 7, [61200, 84000, 90000, 96000, 96000])
    _years(p, 9, sales)
    for row, base in ((11, 98), (12, 9), (13, 11), (14, 36), (15, 154), (16, 121),
                      (19, 18), (20, 6), (21, 0.8), (22, 0.4), (23, 6), (24, 8.3), (25, 39.5),
                      (29, 10), (30, 30), (32, 1.0), (33, 6.5), (34, 26.7), (35, 6.7),
                      (38, 30)):
        _years(p, row, [round(base * (1 + 0.1 * i), 2) for i in range(5)])
    _years(p, 17, [0.44, 0.45, 0.46, 0.46, 0.47])
    _years(p, 27, ebitda)
    _years(p, 28, [e / s for e, s in zip(ebitda, sales)])
    _years(p, 31, TL_INTEREST)
    _years(p, 36, PAT)
    _years(p, 37, [x / s for x, s in zip(PAT, sales)])

    r = wb["Ratios & DSCR"]
    for row, base in ((6, 0.15), (8, 0.07), (9, 0.18), (10, 0.25), (12, 1.6), (13, 1.9),
                      (14, 4.2), (16, 1.5), (20, 45), (23, 36), (24, 18)):
        _years(r, row, [round(base * (1 + 0.05 * i), 3) for i in range(5)])
    _years(r, 25, dscr)
    r["C26"] = sum(dscr) / len(dscr)

    pc = wb["Project Cost"]
    pc["C14"], pc["C21"] = 80, 0

    bs, cf = wb["Balance Sheet"], wb["Cash Flow"]
    six = ("C", "D", "E", "F", "G", "H")
    for row in range(6, 21):
        _years(bs, row, [round(row * (1 + 0.1 * i), 2) for i in range(6)], six)
        _years(cf, row, [round(row * (1 + 0.1 * i), 2) for i in range(6)], six)
    _years(bs, 9, [55.0] + TL_CLOSING, six)
    _years(cf, 15, [0.0] + PRINCIPAL, six)
    bs["A16"], bs["A17"] = "Raw-material stock", "Finished-goods stock"
    _years(bs, 21, [0] * 6, six)
    _years(cf, 21, [0] * 6, six)

    wc = wb["Working Capital"]
    wc["A20"] = "Sundry creditors"
    for row in (18, 20, 22, 23, 24, 27):
        _years(wc, row, [round(row * (1 + 0.1 * i), 2) for i in range(5)])
    be = wb["Break-even & Sensitivity"]
    for row in (6, 7, 9, 10, 13):
        _years(be, row, [round(row * 10 * (1 + 0.1 * i), 2) for i in range(5)])
    _years(be, 8, [0.44] * 5)
    _years(be, 11, [0.62, 0.50, 0.45, 0.41, 0.40])
    _years(be, 12, [0.38, 0.50, 0.55, 0.59, 0.60])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filled_workbook(values=INPUTS) -> bytes:
    from services.template_fill_service import fill_template
    return fill_template("feasibility_study", FEAS_ID, cells(values))


def model_from(recalc: bytes) -> dict:
    return {"narrative": {}, "kpis": [], "sheets": [],
            "financial_summary": frep.financial_summary(recalc),
            "statement_tables": frep.statement_tables(recalc),
            "key_assumptions": frep.key_assumptions(recalc)}


PROJECT = {"title": "Packaging unit", "industry": "Manufacturing", "currency": "INR",
           "country": "India", "location": "Pune", "promoter_name": "A. Promoter",
           "project_cost": 8000000, "loan_amount": 5500000, "own_contribution": 2500000}


@unittest.skipUnless(HAVE_TEMPLATE, "feasibility workbook not installed")
class FeasibilityReadersTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.recalc = fake_recalc(filled_workbook())

    def test_summary_is_in_rupees(self):
        s = frep.financial_summary(self.recalc)
        self.assertEqual(s["series"]["Net Sales / Revenue"][0], 275.4 * LAKH)
        self.assertEqual(s["cards"]["revenue_y1"], 275.4 * LAKH)
        self.assertEqual(s["cards"]["pat_y5"], 72 * LAKH)
        self.assertAlmostEqual(s["cards"]["avg_dscr"], sum(DSCR) / 5)
        self.assertEqual(s["ratios"]["DSCR"], DSCR)
        self.assertAlmostEqual(s["ratios"]["EBITDA Margin"][0], 40 / 275.4)

    def test_verdict_uses_the_same_tiers(self):
        k = frep.real_kpis(self.recalc)
        self.assertFalse(k["scale_flag"])
        self.assertEqual(k["min_dscr"], 1.40)
        self.assertEqual(verdict_tier(k)[0], "STRONG")

    def test_an_out_of_range_margin_is_flagged(self):
        rec = fake_recalc(filled_workbook(), ebitda=[200, 250, 260, 270, 280])
        self.assertEqual(verdict_tier(frep.real_kpis(rec))[0], "REVIEW REQUIRED")

    def test_weak_coverage_is_below_norm(self):
        rec = fake_recalc(filled_workbook(), dscr=[0.8, 0.9, 1.0, 1.1, 1.2])
        self.assertEqual(verdict_tier(frep.real_kpis(rec))[0], "BELOW NORM")

    def test_no_loan_means_no_dscr_and_no_repayment_table(self):
        values = dict(INPUTS, term_loan=0, promoter_equity=80)
        rec = fake_recalc(filled_workbook(values))
        self.assertNotIn("DSCR", frep.financial_summary(rec)["ratios"])
        self.assertNotIn("avg_dscr", frep.real_kpis(rec))
        self.assertIsNone(verdict_tier(frep.real_kpis(rec)))
        titles = [t["title"] for t in frep.statement_tables(rec)]
        self.assertNotIn("Term Loan Repayment Schedule", titles)

    def test_tables_carry_rupees_and_the_labels_the_builders_read(self):
        tables = {t["title"]: t for t in frep.statement_tables(self.recalc)}
        self.assertIn("Consolidated Five-Year Summary", tables)
        self.assertIn("Term Loan Repayment Schedule", tables)
        rows = {r[0]: r[1] for r in tables["Consolidated Five-Year Summary"]["rows"]}
        self.assertEqual(rows["Net Sales (net of GST)"][0], 275.4 * LAKH)
        self.assertEqual(rows["Average capacity utilisation (%)"][0], 51.0)
        self.assertAlmostEqual(rows["EBITDA margin"][0], 40 / 275.4)     # Word shows it as %
        self.assertIn("Output / sales volume — Units (units)", rows)
        repay = {r[0]: r[1] for r in tables["Term Loan Repayment Schedule"]["rows"]}
        self.assertEqual(repay["Principal repayment"], [v * LAKH for v in PRINCIPAL])
        self.assertEqual(repay["Closing balance"], [v * LAKH for v in TL_CLOSING])
        for t in tables.values():
            self.assertTrue(t["key"].startswith("feas_"))     # no CMA commentary on these

    def test_assumptions_in_rupees(self):
        rows = dict(frep.key_assumptions(self.recalc))
        self.assertEqual(rows["TOTAL COST OF PROJECT"], "8,000,000")
        self.assertEqual(rows["Term Loan applied"], "5,500,000")
        self.assertEqual(rows["Promoter's Capital / Equity"], "2,500,000")
        self.assertEqual(rows["Interest rate on Term Loan"], "10.50%")
        self.assertEqual(rows["Plant, machinery & equipment — cost"], "4,000,000")
        self.assertEqual(rows["Direct labour (₹ a month)"], "300,000")
        self.assertEqual(rows["Selling price (₹ per Units)"], "450")
        self.assertNotIn("Land — cost", rows)                  # zero lines are left out

    def test_the_short_report_reads_them(self):
        from services import short_report as sr
        f = sr.short_facts(model_from(self.recalc), PROJECT)
        self.assertEqual(f["project_cost"], 8000000)
        self.assertEqual(f["loan"], 5500000)
        self.assertEqual(f["own"], 2500000)
        self.assertAlmostEqual(f["rate"], 0.105)
        self.assertEqual(f["tenure_months"], 60)
        self.assertEqual(f["moratorium_months"], 6)
        self.assertTrue(f["capex_ties"])
        self.assertAlmostEqual(f["capex_total"], 8000000)
        self.assertEqual(f["revenue"][0], 275.4 * LAKH)
        self.assertEqual(f["pat"][4], 72 * LAKH)
        self.assertEqual(f["dscr"], DSCR)
        self.assertEqual(f["interest_years"], [v * LAKH for v in TL_INTEREST])
        self.assertTrue(f["volume"][0].startswith("Output / sales volume"))
        self.assertTrue(f["costs"])
        self.assertIsNotNone(f["breakeven"])

    def test_short_and_long_documents_build_with_crore_figures(self):
        from docx import Document
        from services import short_report as sr
        from services.word_builder import build_word, _inr
        model = model_from(self.recalc)
        self.assertTrue(sr.build_short_excel(model, PROJECT))
        for data in (sr.build_short_word(model, PROJECT),
                     build_word(model, "feasibility_study", dict(PROJECT, no_cover_art=True))):
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
            self.assertIn(_inr(275.4 * LAKH), text)

    def test_checks_pass_on_a_tallying_workbook(self):
        checks = frep.consistency_checks(self.recalc)
        self.assertEqual(len(checks), 3)
        self.assertTrue(all(c["ok"] for c in checks), checks)

    def test_a_cma_workbook_is_left_to_the_generic_readers(self):
        from template_config import template_path
        path = template_path("bank_loan", "bank_loan_cma")
        with open(path, "rb") as fh:
            data = fh.read()
        self.assertEqual(frep.financial_summary(data), {})
        self.assertEqual(frep.statement_tables(data), [])
        self.assertEqual(frep.key_assumptions(data), [])
        self.assertEqual(frep.real_kpis(data), {})
        self.assertEqual(frep.consistency_checks(data), [])


@unittest.skipUnless(HAVE_TEMPLATE, "feasibility workbook not installed")
class FeasibilityGenerationTests(unittest.TestCase):
    """_run_generation (workbook-only) and the three downloads, on SQLite."""

    def setUp(self):
        from sqlalchemy import create_engine
        from sqlalchemy.orm import sessionmaker
        from sqlalchemy.pool import StaticPool
        from database import Base
        from models.project_model import Project
        from models.user_model import User
        from routers import generation_router as g
        import agents.swot_agent
        import agents.business_model_agent
        import agents.references_agent
        from services import claude_service, excel_cache

        self.g = g
        self.engine = create_engine("sqlite://", poolclass=StaticPool,
                                    connect_args={"check_same_thread": False})
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine, autoflush=False)()
        user = User(email="feas@example.com", hashed_password="x")
        self.db.add(user)
        self.db.commit()
        self.project = Project(title="Packaging unit", user_id=user.id, industry="Manufacturing",
                               purpose="feasibility_study", currency="INR", country="India",
                               financial_format="standard", report_format="long",
                               project_cost=8000000, loan_amount=5500000,
                               own_contribution=2500000)
        self.db.add(self.project)
        self.db.commit()

        self.llm_calls = []

        def tripwire(*a, **k):
            self.llm_calls.append(a[1] if len(a) > 1 else k.get("model"))
            raise RuntimeError("model call attempted in a test")

        # The AI fills the inputs in rupees here, as a model that thinks in rupees would;
        # the Feasibility guard has to bring them back to lakh.
        ai_cells = cells(dict(INPUTS, term_loan=5500000, promoter_equity=25,
                              direct_labour_per_month=300000))
        self.saved = []
        self.cache_dir = tempfile.mkdtemp(prefix="feas_test_cache_")
        for obj, name, value in (
                (claude_service, "_create", tripwire),
                (g, "_build_agent_context", lambda *a, **k: ""),
                (g, "generate_template_inputs", lambda *a, **k: dict(ai_cells)),
                (g, "libreoffice_available", lambda: True),
                (g, "recalculate_xlsx", lambda data, *a, **k: fake_recalc(data)),
                (agents.swot_agent, "swot_agent", lambda **k: ""),
                (g, "swot_agent", lambda **k: ""),
                (agents.business_model_agent, "business_model_agent", lambda *a, **k: None),
                (agents.business_model_agent, "exec_summary_agent", lambda *a, **k: None),
                (agents.references_agent, "references_agent", lambda *a, **k: None),
                (excel_cache, "_DIR", self.cache_dir)):
            self.saved.append((obj, name, getattr(obj, name)))
            setattr(obj, name, value)

    def tearDown(self):
        for obj, name, value in reversed(self.saved):
            setattr(obj, name, value)
        self.db.close()
        self.engine.dispose()
        shutil.rmtree(self.cache_dir, ignore_errors=True)

    def generate(self):
        self.g._run_generation(self.db, self.project, self.g.GenerateRequest(excel_only=True))
        self.db.refresh(self.project)
        return json.loads(self.project.report.financial_model)

    def test_generation_stores_a_rupee_model_from_the_feasibility_workbook(self):
        model = self.generate()
        self.assertEqual(self.llm_calls, [])
        self.assertEqual(model["financial_summary"]["cards"]["revenue_y1"], 275.4 * LAKH)
        titles = [t["title"] for t in model["statement_tables"]]
        self.assertIn("Consolidated Five-Year Summary", titles)
        rows = dict(model["key_assumptions"])
        # The guard turned the AI's rupees into lakh before the workbook was filled.
        self.assertEqual(rows["Term Loan applied"], "5,500,000")
        self.assertEqual(rows["Direct labour (₹ a month)"], "300,000")
        names = [c["name"] for c in model["consistency_checks"]]
        self.assertIn("Balance sheet tallies (assets = liabilities)", names)

    def test_downloads_build_without_a_model_call(self):
        from docx import Document
        from services.word_builder import _inr
        self.generate()
        for fmt in ("long", "short"):
            with self.subTest(report_format=fmt):
                self.project.report_format = fmt
                self.db.commit()
                data, _ = self.g._build_word_report(self.project, self.db, cover_art=False)
                doc = Document(io.BytesIO(data))
                text = "\n".join(c.text for t in doc.tables for row in t.rows
                                 for c in row.cells)
                text += "\n".join(p.text for p in doc.paragraphs)
                self.assertIn(_inr(275.4 * LAKH), text)
                xlsx, fname, _ = self.g._build_excel_file(self.db, self.project)
                self.assertTrue(xlsx)
        self.assertEqual(self.llm_calls, [])


if __name__ == "__main__":
    unittest.main()

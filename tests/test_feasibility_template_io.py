import os
import shutil
import sys
import tempfile
import unittest

BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BACKEND)

from openpyxl import load_workbook  # noqa: E402
from services import feasibility_template_io as t  # noqa: E402

HAVE_TEMPLATE = os.path.isfile(t.TEMPLATE_PATH)
HAVE_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def sample(**over):
    v = {
        "industry": "Food & Restaurant", "start_year": 2026,
        "capacity_per_month": 6000, "working_days_per_month": 26,
        "selling_price_per_unit": 450, "direct_cost_per_unit": 160,
        "util_y2": 0.65, "util_y3": 0.75, "util_y4": 0.80, "util_y5": 0.85,
        "price_escalation": 0.05, "direct_cost_escalation": 0.05, "utility_escalation": 0.04,
        "wage_escalation": 0.07, "overhead_escalation": 0.05,
        "utilities_per_unit": 15, "consumables_per_unit": 10, "packing_per_unit": 8,
        "direct_labour_per_month": 4.5, "admin_salaries_per_month": 2.0,
        "fixed_utility_per_month": 0.5, "repairs_pct_of_equipment": 0.02,
        "insurance_pct_of_assets": 0.004, "other_admin_per_month": 1.0,
        "selling_dist_pct_of_sales": 0.03,
        "term_loan": 80, "promoter_equity": 60, "term_loan_interest": 0.105,
        "term_loan_moratorium_months": 6, "term_loan_repayment_months": 60, "cc_interest": 0.11,
        "stock_days": 10, "finished_stock_days": 5, "debtor_days": 3, "creditor_days": 15,
        "wc_margin": 0.25, "dep_building": 0.0334, "dep_equipment": 0.0633, "dep_electrical": 0.0633,
        "dep_utilities": 0.0633, "dep_misc": 0.1, "preop_amortisation": 0.2,
        "tax_rate": 0.2517, "gst_rate": 0.05,
        "pc_land": 0, "pc_building": 40, "pc_equipment": 60, "pc_electrical": 8,
        "pc_utilities": 5, "pc_misc_assets": 6, "pc_contingency": 4, "pc_preoperative": 7,
        "pc_wc_margin": 10,
    }
    for m in range(1, 13):
        v[f"util_y1_m{m}"] = min(0.30 + 0.03 * (m - 1), 0.60)
    v.update(over)
    return v


class ValidationTests(unittest.TestCase):
    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_good_sample_passes(self):
        t.validate(sample())

    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_rejects_whole_percent(self):
        with self.assertRaises(ValueError):
            t.validate(sample(tax_rate=25))

    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_rejects_unknown_key_negative_and_bad_industry(self):
        for bad in ({"oops": 1}, {"term_loan": -5}, {"industry": "Space Mining"},
                    {"capacity_per_month": "lots"}, {"term_loan_moratorium_months": 6.5}):
            with self.assertRaises(ValueError, msg=str(bad)):
                t.validate(sample(**bad))

    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_missing_inputs_are_reported(self):
        v = sample(); del v["term_loan"]
        with self.assertRaises(ValueError):
            t.validate(v)
        self.assertEqual(t.missing_inputs(v), ["term_loan"])


class WorkbookTests(unittest.TestCase):
    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_every_input_cell_exists_and_is_not_a_formula(self):
        wb = load_workbook(t.TEMPLATE_PATH)
        for key, (sheet, cell, _) in t.INPUTS.items():
            val = wb[sheet][cell].value
            self.assertFalse(isinstance(val, str) and val.startswith("="), f"{key} is a formula")

    @unittest.skipUnless(HAVE_TEMPLATE, "template not found")
    def test_fill_writes_the_right_cells_and_keeps_formulas(self):
        with tempfile.TemporaryDirectory() as d:
            out = t.fill_inputs(sample(), os.path.join(d, "x.xlsx"))
            wb = load_workbook(out)
            self.assertEqual(wb["Assumptions"]["B5"].value, 6000)
            self.assertEqual(wb["Project Cost"]["C6"].value, 40)
            self.assertTrue(str(wb["Project Cost"]["C14"].value).startswith("="))
            self.assertTrue(str(wb["Profitability (Annual)"]["C9"].value).startswith("="))

    @unittest.skipUnless(HAVE_TEMPLATE and HAVE_SOFFICE, "needs template and LibreOffice")
    def test_end_to_end_numbers_reconcile_for_every_industry(self):
        with tempfile.TemporaryDirectory() as d:
            results = {}
            for ind in t.industries():
                p = t.fill_inputs(sample(industry=ind), os.path.join(d, "a.xlsx"))
                t.recalculate(p)
                o = t.read_outputs(p)
                self.assertEqual(o["total_project_cost"], 140)
                self.assertAlmostEqual(o["finance_check"], 0, places=6)
                self.assertTrue(all(abs(x) < 1e-6 for x in o["balance_sheet_check"]), ind)
                self.assertTrue(all(abs(x) < 1e-6 for x in o["cash_flow_check"]), ind)
                self.assertTrue(all(x > 0 for x in o["net_sales"]))
                self.assertEqual(len(o["dscr"]), 5)
                results[ind] = o["net_sales"]
            # the industry only changes labels, never the arithmetic
            first = next(iter(results.values()))
            for ind, s in results.items():
                self.assertEqual(s, first, ind)

    @unittest.skipUnless(HAVE_TEMPLATE and HAVE_SOFFICE, "needs template and LibreOffice")
    def test_year_headings_follow_start_year(self):
        with tempfile.TemporaryDirectory() as d:
            p = t.fill_inputs(sample(start_year=2028), os.path.join(d, "y.xlsx"))
            t.recalculate(p)
            wb = load_workbook(p, data_only=True)
            self.assertEqual(wb["Term Loan"]["C3"].value, "YEAR 1   •   FY2028-29")
            self.assertEqual(wb["Working Capital"]["C4"].value, "Yr1 FY29")
            self.assertEqual(wb["Term Loan"]["C4"].value.year, 2028)


if __name__ == "__main__":
    unittest.main()
